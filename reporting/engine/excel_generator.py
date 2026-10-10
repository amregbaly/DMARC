import os
from typing import Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenpyxlImage

class DmarcExcelGenerator:
    def __init__(self, template_config: Dict[str, Any]):
        self.template = template_config
        self.colors = self.template.get("brand_colors", {})
        
        # Color helper: strip '#' for openpyxl ARGB hex
        def clean_hex(color_str: str, default: str) -> str:
            c = (color_str or default).lstrip("#").upper()
            return "FF" + c if len(c) == 6 else c

        self.primary_hex = clean_hex(self.colors.get("primary"), "0F172A")
        self.header_hex = clean_hex(self.colors.get("table_header"), "1E293B")
        self.header_text_hex = clean_hex(self.colors.get("table_header_text"), "FFFFFF")
        self.pass_hex = clean_hex(self.colors.get("pass_color"), "10B981")
        self.fail_hex = clean_hex(self.colors.get("fail_color"), "EF4444")
        self.warn_hex = clean_hex(self.colors.get("warn_color"), "F59E0B")
        self.bg_light_hex = clean_hex(self.colors.get("background_light"), "F8FAFC")

    def _apply_header_style(self, ws, row_idx: int, col_count: int):
        fill = PatternFill(start_color=self.header_hex, end_color=self.header_hex, fill_type="solid")
        font = Font(name="Calibri", size=11, bold=True, color=self.header_text_hex)
        border = Border(bottom=Side(style="medium", color="000000"))
        
        for col_idx in range(1, col_count + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = border

    def _auto_fit_columns(self, ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    def render_excel(self, data: Dict[str, Any], output_path: str, domain: str = "all") -> str:
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)
        
        summary = data["summary"]
        matrix = data["auth_matrix"]

        # ==========================================
        # 1. SUMMARY SHEET
        # ==========================================
        ws_sum = wb.create_sheet(title="Executive Summary")
        ws_sum.views.sheetView[0].showGridLines = True
        
        # Insert Logo if present
        logo_rel = self.template.get("logo", "templates/assets/sentinel_logo.png")
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        logo_full = os.path.join(base_dir, logo_rel) if not os.path.isabs(logo_rel) else logo_rel
        if os.path.exists(logo_full) and logo_full.endswith(".png"):
            try:
                img = OpenpyxlImage(logo_full)
                img.width = 180
                img.height = 45
                ws_sum.add_image(img, "A1")
            except Exception as e:
                print("Could not attach Excel image:", e)

        # Title Block
        prep_for = self.template.get('prepared_for', 'Enterprise Security Team')
        prep_by = self.template.get('prepared_by', 'DMARC Sentinel')
        ws_sum.cell(row=5, column=1, value=f"Scope: {domain.upper()} | Client: {prep_for} | Author: {prep_by}").font = Font(size=10, italic=True, color="64748B")

        # KPI Block
        ws_sum.cell(row=7, column=1, value="METRIC").font = Font(bold=True)
        ws_sum.cell(row=7, column=2, value="VALUE").font = Font(bold=True)
        self._apply_header_style(ws_sum, 7, 2)

        kpis = [
            ("Total Messages Processed", summary["total_messages"], "#,##0"),
            ("Compliant Messages (Passed DMARC)", summary["passed_messages"], "#,##0"),
            ("Non-Compliant / Threat Messages", summary["failed_messages"], "#,##0"),
            ("Overall DMARC Pass Rate", summary["pass_percentage"] / 100.0, "0.0%"),
            ("Delivered Messages (none)", summary["disposition"]["none"], "#,##0"),
            ("Quarantined Messages (spam)", summary["disposition"]["quarantine"], "#,##0"),
            ("Rejected Messages (blocked)", summary["disposition"]["reject"], "#,##0"),
        ]

        r = 8
        for label, val, num_fmt in kpis:
            c1 = ws_sum.cell(row=r, column=1, value=label)
            c2 = ws_sum.cell(row=r, column=2, value=val)
            c2.number_format = num_fmt
            c2.alignment = Alignment(horizontal="right")
            if "Pass Rate" in label:
                c2.font = Font(bold=True, color="FF10B981" if val >= 0.9 else "FFEF4444")
            r += 1

        # Domain Readiness Table
        r += 2
        ws_sum.cell(row=r, column=1, value="Domain Policy Readiness Assessment").font = Font(size=12, bold=True, color=self.primary_hex)
        r += 1
        headers_dom = ["Monitored Domain", "Message Volume", "Pass Rate", "Current Policy", "Enforcement Recommendation"]
        for ci, h in enumerate(headers_dom, 1):
            ws_sum.cell(row=r, column=ci, value=h)
        self._apply_header_style(ws_sum, r, len(headers_dom))
        
        dom_start_row = r + 1
        r += 1
        for d in data["domains"]:
            ws_sum.cell(row=r, column=1, value=d["domain"])
            c_vol = ws_sum.cell(row=r, column=2, value=d["volume"])
            c_vol.number_format = "#,##0"
            c_vol.alignment = Alignment(horizontal="right")
            c_rate = ws_sum.cell(row=r, column=3, value=d["pass_rate"] / 100.0)
            c_rate.number_format = "0.0%"
            c_rate.alignment = Alignment(horizontal="right")
            ws_sum.cell(row=r, column=4, value=d["current_policy"].upper())
            ws_sum.cell(row=r, column=5, value=d["readiness"]).font = Font(bold=True)
            r += 1

        self._auto_fit_columns(ws_sum)

        # ==========================================
        # 2. AUTHENTICATION SHEET
        # ==========================================
        ws_auth = wb.create_sheet(title="Authentication & Alignment")
        ws_auth.views.sheetView[0].showGridLines = True
        
        ws_auth.cell(row=1, column=1, value="SPF x DKIM Result Alignment Matrix").font = Font(size=13, bold=True, color=self.primary_hex)
        headers_auth = ["Alignment Result State", "SPF Aligned", "DKIM Aligned", "Message Volume", "Percentage", "Security Implication"]
        for ci, h in enumerate(headers_auth, 1):
            ws_auth.cell(row=3, column=ci, value=h)
        self._apply_header_style(ws_auth, 3, len(headers_auth))

        tot = summary["total_messages"] or 1
        matrix_rows = [
            ("Fully Compliant", "YES", "YES", matrix["pass_pass"], matrix["pass_pass"] / tot, "Legitimate corporate mail passing both SPF & DKIM"),
            ("SPF Pass / DKIM Fail", "YES", "NO", matrix["pass_fail"], matrix["pass_fail"] / tot, "SPF aligned, DKIM signature failed or missing"),
            ("SPF Fail / DKIM Pass", "NO", "YES", matrix["fail_pass"], matrix["fail_pass"] / tot, "DKIM aligned, SPF broken by mailing list / forwarder"),
            ("Total Spoofing (Fail/Fail)", "NO", "NO", matrix["fail_fail"], matrix["fail_fail"] / tot, "Unauthorized spoofing / phishing attempts"),
        ]

        r = 4
        for label, spf, dkim, vol, pct, note in matrix_rows:
            ws_auth.cell(row=r, column=1, value=label).font = Font(bold=True)
            ws_auth.cell(row=r, column=2, value=spf)
            ws_auth.cell(row=r, column=3, value=dkim)
            cv = ws_auth.cell(row=r, column=4, value=vol)
            cv.number_format = "#,##0"
            cv.alignment = Alignment(horizontal="right")
            cp = ws_auth.cell(row=r, column=5, value=pct)
            cp.number_format = "0.0%"
            cp.alignment = Alignment(horizontal="right")
            ws_auth.cell(row=r, column=6, value=note)
            r += 1

        # Selectors Table
        r += 2
        ws_auth.cell(row=r, column=1, value="DKIM Selectors Performance").font = Font(size=13, bold=True, color=self.primary_hex)
        r += 1
        headers_sel = ["DKIM Selector", "Message Volume", "Alignment Pass Rate"]
        for ci, h in enumerate(headers_sel, 1):
            ws_auth.cell(row=r, column=ci, value=h)
        self._apply_header_style(ws_auth, r, len(headers_sel))
        r += 1

        for s in data["selectors"]:
            ws_auth.cell(row=r, column=1, value=s["selector"])
            cv = ws_auth.cell(row=r, column=2, value=s["volume"])
            cv.number_format = "#,##0"
            cv.alignment = Alignment(horizontal="right")
            cp = ws_auth.cell(row=r, column=3, value=s["pass_rate"] / 100.0)
            cp.number_format = "0.0%"
            cp.alignment = Alignment(horizontal="right")
            r += 1

        self._auto_fit_columns(ws_auth)

        # ==========================================
        # 3. SOURCES SHEET
        # ==========================================
        ws_src = wb.create_sheet(title="Sender Inventory")
        ws_src.views.sheetView[0].showGridLines = True
        
        headers_send = ["Sender Base Domain", "Provider Name", "Total Volume", "Passed Volume", "Pass Rate", "Classification"]
        for ci, h in enumerate(headers_send, 1):
            ws_src.cell(row=1, column=ci, value=h)
        self._apply_header_style(ws_src, 1, len(headers_send))
        ws_src.freeze_panes = "A2"

        r = 2
        for s in data["senders"]:
            ws_src.cell(row=r, column=1, value=s["base_domain"]).font = Font(bold=True)
            ws_src.cell(row=r, column=2, value=s["name"])
            c_tot = ws_src.cell(row=r, column=3, value=s["volume"])
            c_tot.number_format = "#,##0"
            c_pass = ws_src.cell(row=r, column=4, value=s["passed"])
            c_pass.number_format = "#,##0"
            c_rate = ws_src.cell(row=r, column=5, value=s["pass_rate"] / 100.0)
            c_rate.number_format = "0.0%"
            c_rate.alignment = Alignment(horizontal="right")
            
            c_st = ws_src.cell(row=r, column=6, value=s["status"])
            if s["status"] == "Legitimate":
                c_st.font = Font(color="FF10B981", bold=True)
            elif s["status"] == "Misconfigured":
                c_st.font = Font(color="FFF59E0B", bold=True)
            else:
                c_st.font = Font(color="FFEF4444", bold=True)
            r += 1

        self._auto_fit_columns(ws_src)

        # ==========================================
        # 4. SECURITY SHEET
        # ==========================================
        ws_sec = wb.create_sheet(title="Security & Spoofing")
        ws_sec.views.sheetView[0].showGridLines = True
        
        headers_sec = ["Attacking Source IP", "Reverse DNS (PTR)", "Origin Country", "Autonomous System (ISP)", "Failing Volume (Blocked)"]
        for ci, h in enumerate(headers_sec, 1):
            ws_sec.cell(row=1, column=ci, value=h)
        self._apply_header_style(ws_sec, 1, len(headers_sec))
        ws_sec.freeze_panes = "A2"

        r = 2
        for t in data["threats"]:
            ws_sec.cell(row=r, column=1, value=t["ip"]).font = Font(name="Courier New", bold=True)
            ws_sec.cell(row=r, column=2, value=t["ptr"]).font = Font(name="Courier New", size=9)
            ws_sec.cell(row=r, column=3, value=t["country"])
            ws_sec.cell(row=r, column=4, value=t["as_name"])
            cv = ws_sec.cell(row=r, column=5, value=t["volume"])
            cv.number_format = "#,##0"
            cv.font = Font(color="FFEF4444", bold=True)
            cv.alignment = Alignment(horizontal="right")
            r += 1

        self._auto_fit_columns(ws_sec)

        # ==========================================
        # 5. RAW DATA SHEET
        # ==========================================
        ws_raw = wb.create_sheet(title="Raw Aggregate Records")
        ws_raw.views.sheetView[0].showGridLines = True
        
        headers_raw = ["Report Date", "Header From", "Source IP", "Sender Domain", "Country", "Messages", "Disposition", "DMARC", "SPF Aligned", "DKIM Aligned"]
        for ci, h in enumerate(headers_raw, 1):
            ws_raw.cell(row=1, column=ci, value=h)
        self._apply_header_style(ws_raw, 1, len(headers_raw))
        ws_raw.freeze_panes = "A2"

        r = 2
        for rec in data["raw_records"]:
            ws_raw.cell(row=r, column=1, value=rec["date"])
            ws_raw.cell(row=r, column=2, value=rec["header_from"])
            ws_raw.cell(row=r, column=3, value=rec["source_ip"]).font = Font(name="Courier New")
            ws_raw.cell(row=r, column=4, value=rec["sender"])
            ws_raw.cell(row=r, column=5, value=rec["country"])
            cv = ws_raw.cell(row=r, column=6, value=rec["count"])
            cv.number_format = "#,##0"
            cv.alignment = Alignment(horizontal="right")
            ws_raw.cell(row=r, column=7, value=rec["disposition"])
            
            cdm = ws_raw.cell(row=r, column=8, value=rec["dmarc"])
            cdm.font = Font(color="FF10B981" if rec["dmarc"] == "PASS" else "FFEF4444", bold=True)
            
            ws_raw.cell(row=r, column=9, value=rec["spf_aligned"])
            ws_raw.cell(row=r, column=10, value=rec["dkim_aligned"])
            r += 1

        self._auto_fit_columns(ws_raw)

        # Save workbook
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        wb.save(output_path)
        return output_path
