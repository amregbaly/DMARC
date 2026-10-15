import os
import sys
import json
import base64
import argparse
from datetime import datetime
from typing import Dict, Any, List

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright
try:
    from engine.logo_resolver import get_logo_data_uri, get_logo_openpyxl_image
except ImportError:
    try:
        from reporting.engine.logo_resolver import get_logo_data_uri, get_logo_openpyxl_image
    except ImportError:
        from logo_resolver import get_logo_data_uri, get_logo_openpyxl_image

# -------------------------------------------------------------
# Color Palette Constants (Executive Light)
# -------------------------------------------------------------
HEX_PRIMARY = "0F172A"       # Dark Slate / Navy
HEX_ACCENT = "0284C7"        # Azure / Sky Blue
HEX_HEADER_BG = "1E293B"     # Slate Header
HEX_HEADER_TEXT = "FFFFFF"   # White
HEX_ZEBRA = "F8FAFC"         # Light Tint
HEX_BORDER = "CBD5E1"        # Slate 300 Border
HEX_PASS = "10B981"          # Emerald
HEX_FAIL = "EF4444"          # Rose
HEX_WARN = "F59E0B"          # Amber
HEX_MUTED = "64748B"         # Muted Slate

def clean_hex(color_str: str) -> str:
    c = (color_str or "").lstrip("#").upper()
    return "FF" + c if len(c) == 6 else c

# -------------------------------------------------------------
# TAB EXPORTER CLASS
# -------------------------------------------------------------
class TabExporter:
    def __init__(self, data: Dict[str, Any], output_dir: str):
        self.data = data
        self.output_dir = output_dir
        self.tab_id = data.get("tab_id", "aggregate_view")
        self.tab_title = data.get("tab_title", "Aggregate View Analytics")
        self.domain = data.get("domain", "all")
        self.time_window = data.get("time_window", "Last 30 Days")
        self.summary = data.get("summary", {})
        self.columns = data.get("columns", [])
        self.rows = data.get("rows", [])
        self.stats = data.get("stats", [])
        self.filters = data.get("filters", {})
        self.extra = data.get("extra", {})
        self.client_org = data.get("client_org", "Enterprise Security")
        self.user_logon = data.get("user_logon", "System Administrator")
        self.logo_url = (data.get("logo_url") or data.get("logo") or "").strip()
        self.lang = data.get("language", "en")
        self.now_utc = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")

        os.makedirs(self.output_dir, exist_ok=True)

    def _get_logo_data_uri(self) -> str:
        return get_logo_data_uri(self.logo_url, fallback=True)

    def _get_filename(self, ext: str) -> str:
        safe_domain = "".join([c if c.isalnum() or c in ("-", "_") else "_" for c in self.domain]).strip("_") or "all"
        safe_tab = "".join([c if c.isalnum() or c in ("-", "_") else "_" for c in self.tab_id]).strip("_") or "tab"
        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        return f"dmarc_{safe_tab}_{safe_domain}_{timestamp}.{ext}"

    # =========================================================
    # EXCEL EXPORT GENERATION (.xlsx)
    # =========================================================
    def export_excel(self) -> Dict[str, Any]:
        filename = self._get_filename("xlsx")
        output_path = os.path.join(self.output_dir, filename)

        wb = openpyxl.Workbook()
        ws = wb.active
        # Safe sheet name (max 31 chars, no special characters)
        safe_title = (self.tab_title[:28] if len(self.tab_title) > 28 else self.tab_title)
        for ch in [':', '\\', '/', '?', '*', '[', ']']:
            safe_title = safe_title.replace(ch, ' ')
        ws.title = safe_title.strip() or "Aggregate Data"
        ws.views.sheetView[0].showGridLines = True

        thin_side = Side(style="thin", color="FF" + HEX_BORDER)
        cell_border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
        header_fill = PatternFill(start_color="FF" + HEX_HEADER_BG, end_color="FF" + HEX_HEADER_BG, fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FF" + HEX_HEADER_TEXT)
        zebra_fill = PatternFill(start_color="FF" + HEX_ZEBRA, end_color="FF" + HEX_ZEBRA, fill_type="solid")
        white_fill = PatternFill(start_color="FFFFFFFF", end_color="FFFFFFFF", fill_type="solid")

        # 1. Embed Organization Logo
        img = get_logo_openpyxl_image(self.logo_url, max_width=180, max_height=42, fallback=True)
        start_row = 1
        if img:
            try:
                ws.add_image(img, "A1")
                start_row = 4
            except Exception as e:
                print("Could not attach Excel image:", e)

        # 2. Document Title & Metadata Block
        ws.cell(row=start_row, column=1, value=f"DMARC SENTINEL AGGREGATE AUDIT — {self.tab_title.upper()}").font = Font(
            name="Calibri", size=15, bold=True, color="FF" + HEX_PRIMARY
        )
        ws.cell(row=start_row + 1, column=1, value=f"Domain: {self.domain.upper()} | Period: {self.time_window} | Generated: {self.now_utc}").font = Font(
            name="Calibri", size=10, italic=True, color="FF" + HEX_MUTED
        )
        ws.cell(row=start_row + 2, column=1, value=f"Client Org: {self.client_org} | Operator: {self.user_logon}").font = Font(
            name="Calibri", size=9, color="FF" + HEX_MUTED
        )

        filter_parts = []
        if self.filters.get("search"):
            filter_parts.append(f"Search: '{self.filters['search']}'")
        if self.filters.get("category") and self.filters["category"] != "all":
            filter_parts.append(f"Category: '{self.filters['category']}'")
        if self.filters.get("hostnameSearch"):
            filter_parts.append(f"Hostname: '{self.filters['hostnameSearch']}'")
        if self.filters.get("ipSearch"):
            filter_parts.append(f"IP: '{self.filters['ipSearch']}'")

        next_row = start_row + 3
        if filter_parts:
            ws.cell(row=next_row, column=1, value=f"Active Table Filters: {' | '.join(filter_parts)}").font = Font(
                name="Calibri", size=9, bold=True, color="FF" + HEX_ACCENT
            )
            next_row += 1

        # 3. Executive KPI Summary Cards
        row_idx = next_row + 1
        kpis = [
            ("TOTAL MESSAGES", self.summary.get("totalVolume", 0), "#,##0"),
            ("COMPLIANT (ALIGNED)", self.summary.get("compliantVolume", 0), "#,##0"),
            ("FORWARDED RELAY", self.summary.get("forwardedVolume", 0), "#,##0"),
            ("NON-COMPLIANT / THREATS", self.summary.get("failedVolume", self.summary.get("threatVolume", 0)), "#,##0"),
            ("DMARC COMPLIANCE RATE", (self.summary.get("complianceRate", 100) or 0) / 100.0, "0.0%")
        ]

        for i, (kpi_label, kpi_val, kpi_fmt) in enumerate(kpis, start=1):
            lbl_cell = ws.cell(row=row_idx, column=i, value=kpi_label)
            lbl_cell.font = Font(name="Calibri", size=9, bold=True, color="FF" + HEX_MUTED)
            lbl_cell.fill = zebra_fill
            lbl_cell.alignment = Alignment(horizontal="center", vertical="center")
            lbl_cell.border = cell_border

            val_cell = ws.cell(row=row_idx + 1, column=i, value=kpi_val)
            val_cell.font = Font(name="Calibri", size=13, bold=True, color="FF" + HEX_PRIMARY)
            val_cell.alignment = Alignment(horizontal="center", vertical="center")
            val_cell.number_format = kpi_fmt
            val_cell.border = cell_border

        row_idx += 4

        # 3. Category Stats strip (if present)
        if self.stats:
            ws.cell(row=row_idx, column=1, value="CATEGORY SUMMARY BREAKDOWN").font = Font(name="Calibri", size=11, bold=True, color="FF" + HEX_PRIMARY)
            row_idx += 1
            for stat in self.stats:
                lbl = stat.get("label", "")
                cnt = stat.get("count", 0)
                vol = stat.get("volume", 0)
                txt = f"{lbl}: {cnt:,} sources | {vol:,} messages"
                ws.cell(row=row_idx, column=1, value=txt).font = Font(name="Calibri", size=10, bold=True, color="FF" + HEX_ACCENT)
                row_idx += 1
            row_idx += 1

        # 4. Table Headers
        table_start_row = row_idx
        col_count = len(self.columns)
        if col_count == 0:
            # Fallback if no explicit column defs
            self.columns = [{"key": "name", "label": "Item", "align": "left", "format": "text"}]
            col_count = 1

        for col_idx, col in enumerate(self.columns, start=1):
            cell = ws.cell(row=table_start_row, column=col_idx, value=col.get("label", col.get("key", "")).upper())
            cell.fill = header_fill
            cell.font = header_font
            align_h = col.get("align", "left")
            cell.alignment = Alignment(horizontal=align_h, vertical="center", wrap_text=True)
            cell.border = cell_border

        # 5. Table Rows
        current_row = table_start_row + 1
        for r_i, row_data in enumerate(self.rows):
            fill = zebra_fill if r_i % 2 == 1 else white_fill
            for col_idx, col in enumerate(self.columns, start=1):
                key = col.get("key")
                fmt = col.get("format", "text")
                align_h = col.get("align", "left")
                val = row_data.get(key, "")

                cell = ws.cell(row=current_row, column=col_idx)
                cell.fill = fill
                cell.border = cell_border
                cell.alignment = Alignment(horizontal=align_h, vertical="center")

                # Type conversions & formatting
                if fmt == "number":
                    try:
                        num = int(val) if isinstance(val, (int, float, str)) and str(val).replace('-', '').isdigit() else float(val)
                        cell.value = num
                        cell.number_format = "#,##0"
                    except Exception:
                        cell.value = str(val or 0)
                elif fmt == "percentage":
                    try:
                        p_val = float(str(val).rstrip("%"))
                        cell.value = p_val / 100.0
                        cell.number_format = "0.0%"
                    except Exception:
                        cell.value = str(val or "0%")
                elif fmt == "boolean":
                    cell.value = "PASS" if val else "FAIL"
                    cell.font = Font(name="Calibri", size=10, bold=True, color="FF" + (HEX_PASS if val else HEX_FAIL))
                else:
                    cell.value = str(val) if val is not None else ""
                    if key in ("ip", "ptr", "hostname", "lat", "lon", "countryCode"):
                        cell.font = Font(name="Consolas", size=10)

            current_row += 1

        # Freeze table headers
        ws.freeze_panes = f"A{table_start_row + 1}"

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

        wb.save(output_path)
        return {
            "format": "xlsx",
            "filename": filename,
            "path": output_path,
            "size_bytes": os.path.getsize(output_path)
        }

    # =========================================================
    # PDF EXPORT GENERATION (.pdf)
    # =========================================================
    def export_pdf(self) -> Dict[str, Any]:
        filename = self._get_filename("pdf")
        output_path = os.path.join(self.output_dir, filename)

        logo_uri = self._get_logo_data_uri()
        html = self._generate_html(logo_uri)

        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page()
            page.set_content(html, wait_until="networkidle")

            page.pdf(
                path=output_path,
                format="A4",
                landscape=True,  # Aggregate tables are wide and benefit significantly from landscape
                print_background=True,
                margin={"top": "12mm", "bottom": "14mm", "left": "12mm", "right": "12mm"},
                display_header_footer=True,
                header_template=f'''
                    <div style="font-family: -apple-system, BlinkMacSystemFont, Segoe UI, Roboto, sans-serif; font-size: 8pt; color: #94a3b8; width: 100%; padding: 0 12mm; display: flex; justify-content: space-between; align-items: center;">
                        <span style="font-weight: 700; color: #0284c7;">DMARC SENTINEL — AGGREGATE SECURITY INTELLIGENCE</span>
                        <span>Scope: <strong>{self.domain}</strong> | {self.tab_title}</span>
                    </div>
                ''',
                footer_template='''
                    <div style="font-family: -apple-system, BlinkMacSystemFont, Segoe UI, Roboto, sans-serif; font-size: 8pt; color: #94a3b8; width: 100%; padding: 0 12mm; display: flex; justify-content: space-between; align-items: center;">
                        <span>CONFIDENTIAL: Authorized Security Operations Use Only</span>
                        <span>Page <span class="pageNumber"></span> of <span class="totalPages"></span></span>
                    </div>
                '''
            )
            browser.close()

        return {
            "format": "pdf",
            "filename": filename,
            "path": output_path,
            "size_bytes": os.path.getsize(output_path)
        }

    # ---------------------------------------------------------
    # HTML Builder for PDF
    # ---------------------------------------------------------
    def _generate_html(self, logo_uri: str) -> str:
        # Format metrics
        tot = self.summary.get("totalVolume", 0)
        comp = self.summary.get("compliantVolume", 0)
        fwd = self.summary.get("forwardedVolume", 0)
        fail = self.summary.get("failedVolume", self.summary.get("threatVolume", 0))
        rate = self.summary.get("complianceRate", 100) or 0

        # Badges for active filters
        filter_html = ""
        filters_list = []
        if self.filters.get("search"):
            filters_list.append(f"Search: <strong>{self.filters['search']}</strong>")
        if self.filters.get("category") and self.filters["category"] != "all":
            filters_list.append(f"Category: <strong>{self.filters['category']}</strong>")
        if self.filters.get("hostnameSearch"):
            filters_list.append(f"Host: <strong>{self.filters['hostnameSearch']}</strong>")
        if self.filters.get("ipSearch"):
            filters_list.append(f"IP: <strong>{self.filters['ipSearch']}</strong>")

        if filters_list:
            filter_html = f'''
            <div class="filter-strip">
                <span class="filter-tag">Active Filter:</span> {' &bull; '.join(filters_list)}
            </div>
            '''

        # Category stats strip if present
        stats_html = ""
        if self.stats:
            stats_pills = []
            for s in self.stats:
                lbl = s.get("label", "")
                cnt = s.get("count", 0)
                vol = s.get("volume", 0)
                color = "#10b981" if "Compatible" in lbl else ("#3b82f6" if "Forwarded" in lbl else "#ef4444")
                stats_pills.append(f'''
                    <div class="stat-pill" style="border-left: 3px solid {color};">
                        <div class="stat-pill-label">{lbl}</div>
                        <div class="stat-pill-val">{cnt:,} <span style="font-size: 8pt; color: #64748b; font-weight: normal;">({vol:,} msgs)</span></div>
                    </div>
                ''')
            stats_html = f'<div class="stats-grid">{"".join(stats_pills)}</div>'

        # Special view: per_result rich cards
        per_result_html = ""
        if self.tab_id == "per_result" and self.extra.get("perResultGroups"):
            groups = self.extra["perResultGroups"]
            cg = groups.get("correctGroup", {})
            ag = groups.get("authenticatedGroup", {})
            ig = groups.get("invalidGroup", {})

            def render_group_card(grp, color, border_color, badge_text):
                items = grp.get("subItems", [])
                items_rows = "".join([
                    f'''<div class="flow-item">
                        <span>{it.get("name", "")}</span>
                        <span class="flow-num">{it.get("count", 0):,} ({it.get("pct", 0)}%)</span>
                    </div>''' for it in items
                ])
                return f'''
                <div class="result-card" style="border-top: 4px solid {color};">
                    <div class="result-badge" style="background: {color}15; color: {color}; border: 1px solid {border_color};">{badge_text}</div>
                    <div class="result-title">{grp.get("title", "")}</div>
                    <div class="result-vol" style="color: {color};">{grp.get("totalVolume", 0):,}</div>
                    <div class="result-pct">{grp.get("percentage", 0)}% of total traffic</div>
                    <div class="flow-items-box">{items_rows}</div>
                </div>
                '''

            per_result_html = f'''
            <div class="result-cards-grid">
                {render_group_card(cg, "#10b981", "#86efac", "DMARC CORRECT / PASS")}
                {render_group_card(ag, "#3b82f6", "#93c5fd", "VALID NOT ALIGNED")}
                {render_group_card(ig, "#ef4444", "#fca5a5", "SPOOFING / REJECTED")}
            </div>
            '''

        # Table rows generation
        table_rows_html = []
        for r_i, row in enumerate(self.rows):
            zebra_cls = "row-alt" if r_i % 2 == 1 else ""
            cells_html = []
            for col in self.columns:
                key = col.get("key")
                fmt = col.get("format", "text")
                align = col.get("align", "left")
                val = row.get(key, "")

                cell_content = ""
                if fmt == "number":
                    try:
                        num = int(val) if isinstance(val, (int, float, str)) and str(val).replace('-', '').isdigit() else float(val)
                        cell_content = f"{num:,}"
                    except Exception:
                        cell_content = str(val or 0)
                    cells_html.append(f'<td style="text-align: {align}; font-family: ui-monospace, monospace; font-weight: 600;">{cell_content}</td>')
                elif fmt == "percentage":
                    p_val = float(str(val).rstrip("%") or 0)
                    color = "#10b981" if p_val >= 95 else ("#f59e0b" if p_val >= 70 else "#ef4444")
                    cells_html.append(f'<td style="text-align: {align}; font-family: ui-monospace, monospace; font-weight: 700; color: {color};">{p_val:.1f}%</td>')
                elif fmt == "boolean":
                    if val:
                        cells_html.append(f'<td style="text-align: {align};"><span class="badge badge-pass">PASS</span></td>')
                    else:
                        cells_html.append(f'<td style="text-align: {align};"><span class="badge badge-fail">FAIL</span></td>')
                elif key == "categoryTag":
                    tag = str(val)
                    if "compatible" in tag.lower():
                        cells_html.append(f'<td style="text-align: {align};"><span class="badge badge-pass">{tag}</span></td>')
                    elif "forwarded" in tag.lower():
                        cells_html.append(f'<td style="text-align: {align};"><span class="badge badge-fwd">{tag}</span></td>')
                    else:
                        cells_html.append(f'<td style="text-align: {align};"><span class="badge badge-fail">{tag}</span></td>')
                elif key in ("ip", "ptr", "hostname"):
                    cells_html.append(f'<td style="text-align: {align}; font-family: ui-monospace, monospace; font-size: 8.5pt; font-weight: 600; color: #0284c7;">{val or "—"}</td>')
                else:
                    cells_html.append(f'<td style="text-align: {align};">{val if val is not None else ""}</td>')

            table_rows_html.append(f'<tr class="{zebra_cls}">{"".join(cells_html)}</tr>')

        headers_html = "".join([
            f'<th style="text-align: {col.get("align", "left")};">{col.get("label", col.get("key", "")).upper()}</th>'
            for col in self.columns
        ])

        html = f'''<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>DMARC Sentinel — {self.tab_title}</title>
<style>
  @page {{
    size: A4 landscape;
    margin: 12mm 12mm 14mm 12mm;
  }}
  * {{
    box-sizing: border-box;
    -webkit-print-color-adjust: exact !important;
    print-color-adjust: exact !important;
  }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
    font-size: 9pt;
    line-height: 1.35;
    color: #0f172a;
    background: #ffffff;
    margin: 0;
    padding: 0;
  }}
  
  /* Header section */
  .doc-header {{
    display: flex;
    justify-content: space-between;
    align-items: center;
    border-bottom: 2px solid #0284c7;
    padding-bottom: 12px;
    margin-bottom: 14px;
  }}
  .brand-block {{
    display: flex;
    align-items: center;
    gap: 12px;
  }}
  .brand-logo {{
    height: 38px;
    object-fit: contain;
  }}
  .brand-title {{
    font-size: 15pt;
    font-weight: 800;
    color: #0f172a;
    letter-spacing: -0.3px;
    line-height: 1.1;
  }}
  .brand-subtitle {{
    font-size: 8.5pt;
    font-weight: 600;
    color: #0284c7;
    text-transform: uppercase;
    letter-spacing: 0.5px;
    margin-top: 2px;
  }}
  .meta-block {{
    text-align: right;
    font-size: 8pt;
    color: #64748b;
  }}
  .meta-badge {{
    display: inline-block;
    padding: 3px 8px;
    background: #f1f5f9;
    border: 1px solid #cbd5e1;
    border-radius: 6px;
    font-weight: 700;
    color: #0f172a;
    margin-bottom: 3px;
  }}

  /* KPI Grid */
  .kpi-grid {{
    display: grid;
    grid-template-columns: repeat(5, 1fr);
    gap: 10px;
    margin-bottom: 14px;
  }}
  .kpi-card {{
    background: #f8fafc;
    border: 1px solid #e2e8f0;
    border-radius: 8px;
    padding: 10px 12px;
  }}
  .kpi-card.kpi-main {{
    background: #f0fdf4;
    border-color: #bbf7d0;
  }}
  .kpi-title {{
    font-size: 7.5pt;
    font-weight: 700;
    color: #64748b;
    text-transform: uppercase;
    letter-spacing: 0.5px;
  }}
  .kpi-value {{
    font-size: 14pt;
    font-weight: 900;
    color: #0f172a;
    font-family: ui-monospace, monospace;
    margin-top: 2px;
    line-height: 1.1;
  }}
  .kpi-sub {{
    font-size: 7pt;
    color: #94a3b8;
    margin-top: 2px;
  }}

  /* Filter Strip */
  .filter-strip {{
    background: #eff6ff;
    border: 1px solid #bfdbfe;
    border-radius: 6px;
    padding: 5px 10px;
    font-size: 8pt;
    color: #1e40af;
    margin-bottom: 12px;
  }}
  .filter-tag {{
    font-weight: 800;
    text-transform: uppercase;
    font-size: 7pt;
    margin-right: 4px;
  }}

  /* Stats Grid */
  .stats-grid {{
    display: grid;
    grid-template-columns: repeat(3, 1fr);
    gap: 10px;
    margin-bottom: 14px;
  }}
  .stat-pill {{
    background: #f8fafc;
    border: 1px solid #e2e8f0;
    border-radius: 6px;
    padding: 8px 12px;
  }}
  .stat-pill-label {{
    font-size: 7.5pt;
    font-weight: 700;
    text-transform: uppercase;
    color: #475569;
  }}
  .stat-pill-val {{
    font-size: 11pt;
    font-weight: 800;
    color: #0f172a;
    font-family: ui-monospace, monospace;
    margin-top: 2px;
  }}

  /* Result Groups */
  .result-cards-grid {{
    display: grid;
    grid-template-columns: repeat(3, 1fr);
    gap: 12px;
    margin-bottom: 14px;
  }}
  .result-card {{
    background: #f8fafc;
    border: 1px solid #e2e8f0;
    border-radius: 8px;
    padding: 12px;
  }}
  .result-badge {{
    display: inline-block;
    padding: 2px 6px;
    border-radius: 4px;
    font-size: 7pt;
    font-weight: 800;
    text-transform: uppercase;
    margin-bottom: 6px;
  }}
  .result-title {{
    font-size: 10pt;
    font-weight: 800;
    color: #0f172a;
  }}
  .result-vol {{
    font-size: 16pt;
    font-weight: 900;
    font-family: ui-monospace, monospace;
    margin-top: 4px;
  }}
  .result-pct {{
    font-size: 8pt;
    font-weight: 600;
    color: #64748b;
    margin-bottom: 8px;
  }}
  .flow-items-box {{
    border-top: 1px solid #e2e8f0;
    padding-top: 6px;
    font-size: 7.5pt;
  }}
  .flow-item {{
    display: flex;
    justify-content: space-between;
    padding: 2px 0;
    border-bottom: 1px dotted #e2e8f0;
  }}
  .flow-num {{
    font-family: ui-monospace, monospace;
    font-weight: 700;
    color: #0f172a;
  }}

  /* Main Table */
  table.data-table {{
    width: 100%;
    border-collapse: collapse;
    font-size: 8pt;
    margin-top: 4px;
  }}
  table.data-table thead {{
    display: table-header-group;
  }}
  table.data-table tr {{
    page-break-inside: avoid;
  }}
  table.data-table th {{
    background: #0f172a;
    color: #ffffff;
    font-size: 7pt;
    font-weight: 800;
    text-transform: uppercase;
    letter-spacing: 0.5px;
    padding: 7px 8px;
    border: 1px solid #0f172a;
  }}
  table.data-table td {{
    padding: 6px 8px;
    border-bottom: 1px solid #e2e8f0;
    border-left: 1px solid #f1f5f9;
    border-right: 1px solid #f1f5f9;
    color: #1e293b;
  }}
  table.data-table tr.row-alt td {{
    background: #f8fafc;
  }}

  /* Badges */
  .badge {{
    display: inline-block;
    padding: 2px 6px;
    border-radius: 4px;
    font-size: 6.5pt;
    font-weight: 800;
    text-transform: uppercase;
  }}
  .badge-pass {{
    background: #dcfce7;
    color: #15803d;
    border: 1px solid #86efac;
  }}
  .badge-fail {{
    background: #fee2e2;
    color: #b91c1c;
    border: 1px solid #fca5a5;
  }}
  .badge-fwd {{
    background: #dbeafe;
    color: #1d4ed8;
    border: 1px solid #93c5fd;
  }}

  /* Footer Note */
  .footer-note {{
    margin-top: 14px;
    font-size: 7pt;
    color: #94a3b8;
    text-align: right;
  }}
</style>
</head>
<body>

<!-- Header -->
<div class="doc-header">
  <div class="brand-block">
    {"<img src='" + logo_uri + "' class='brand-logo' alt='Sentinel' />" if logo_uri else ""}
    <div>
      <div class="brand-title">DMARC Sentinel — {self.tab_title}</div>
      <div class="brand-subtitle">Aggregate Authentication &amp; Delivery Telemetry Audit</div>
    </div>
  </div>
  <div class="meta-block">
    <div><span class="meta-badge">Domain: {self.domain}</span></div>
    <div>Period: <strong>{self.time_window}</strong></div>
    <div>Generated: <strong>{self.now_utc}</strong></div>
    <div>Client: {self.client_org} | User: {self.user_logon}</div>
  </div>
</div>

<!-- KPI Cards -->
<div class="kpi-grid">
  <div class="kpi-card">
    <div class="kpi-title">Total Messages</div>
    <div class="kpi-value">{tot:,}</div>
    <div class="kpi-sub">Aggregated records</div>
  </div>
  <div class="kpi-card kpi-main">
    <div class="kpi-title">DMARC Pass Rate</div>
    <div class="kpi-value" style="color: #10b981;">{rate:.1f}%</div>
    <div class="kpi-sub">{comp:,} aligned</div>
  </div>
  <div class="kpi-card">
    <div class="kpi-title">Compliant Volume</div>
    <div class="kpi-value" style="color: #10b981;">{comp:,}</div>
    <div class="kpi-sub">Passed DMARC check</div>
  </div>
  <div class="kpi-card">
    <div class="kpi-title">Forwarded Traffic</div>
    <div class="kpi-value" style="color: #3b82f6;">{fwd:,}</div>
    <div class="kpi-sub">Indirect relays</div>
  </div>
  <div class="kpi-card">
    <div class="kpi-title">Threat / Failed</div>
    <div class="kpi-value" style="color: #ef4444;">{fail:,}</div>
    <div class="kpi-sub">Unauthenticated traffic</div>
  </div>
</div>

{filter_html}
{stats_html}
{per_result_html}

<!-- Main Data Table -->
<table class="data-table">
  <thead>
    <tr>
      {headers_html}
    </tr>
  </thead>
  <tbody>
    {"".join(table_rows_html) if table_rows_html else '<tr><td colspan="' + str(len(self.columns)) + '" style="text-align:center; padding: 20px; color: #94a3b8;">No records match the current filter selection.</td></tr>'}
  </tbody>
</table>

<div class="footer-note">
  Showing {len(self.rows):,} records &bull; Extracted live from OpenSearch aggregate index cluster &bull; DMARC Sentinel Engine
</div>

</body>
</html>'''
        return html


# -------------------------------------------------------------
# CLI Entrypoint
# -------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(description="Export any Aggregate View tab data to PDF or Excel")
    parser.add_argument("--input", required=True, help="Path to input JSON file containing tab data")
    parser.add_argument("--format", required=True, choices=["pdf", "xlsx"], help="Export format")
    parser.add_argument("--output-dir", required=True, help="Directory to store the exported report")

    args = parser.parse_args()

    if not os.path.exists(args.input):
        print(json.dumps({"status": "error", "message": f"Input file not found: {args.input}"}))
        sys.exit(1)

    try:
        with open(args.input, "r", encoding="utf-8") as f:
            data = json.load(f)

        exporter = TabExporter(data, args.output_dir)

        if args.format == "xlsx":
            res = exporter.export_excel()
        else:
            res = exporter.export_pdf()

        print(json.dumps({"status": "success", **res}))
        sys.exit(0)

    except Exception as e:
        import traceback
        err_msg = traceback.format_exc()
        print(json.dumps({"status": "error", "message": str(e), "traceback": err_msg}))
        sys.exit(2)

if __name__ == "__main__":
    main()
